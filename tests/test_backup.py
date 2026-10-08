"""Backup de Información: ZIP transmitido por partes, firmas como archivos.

Pedido de SESURSA (2026-10-08): "Error 500" al generar el backup. Cloud Run
registraba "Response size was too large" en cada intento (tope de 32 MiB por
respuesta enviada de una pieza) y, en el del año completo, el agotamiento de
los 512 MiB de memoria. Las firmas en base64 iban dos veces (JSON y Excel) y
todo se armaba en memoria. Además la fila de backups_realizados se insertaba
antes de enviar, así que "Último backup hace 1 día" era falso. Esta prueba fija:

- la respuesta no declara Content-Length (va por partes) y el ZIP abre;
- las firmas salen del JSON y del Excel y van en firmas/<formulario>/ como
  imagen idéntica a la original; el registro guarda la ruta; también las
  anidadas (listas de asistencia);
- la fila de backups_realizados sólo existe si el ZIP se transmitió completo.

Usa el arranque común de tests/sekapp_testing.py (Postgres desechable). Correr con:
    monolith/venv/bin/python tests/test_backup.py
"""
import base64
import io
import json
import unittest
import zipfile

from sekapp_testing import A, sql, configurar_app, recrear_base, crear_empresa, crear_usuario, login

ADMIN = 'admin@pruebas.sekapp'
# PNG de 1x1 válido; la firma viaja como data URI, igual que en producción.
PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==')
FIRMA = 'data:image/png;base64,' + base64.b64encode(PNG).decode()
PERIODO = {'start_date': '2000-01-01', 'end_date': '2099-12-31'}


def setUpModule():
    configurar_app()
    recrear_base()
    cid = crear_empresa()
    crear_usuario(ADMIN, 'Admin Pruebas', cid, is_admin=True, is_super_admin=True)
    sql("INSERT INTO reportes_incidentes (cliente_instalacion, estado, user_email, creado_en, firma_responsable) "
        "VALUES ('Instalación B1', 'Abierto', %s, NOW(), %s)", [ADMIN, FIRMA])
    sql("INSERT INTO reportes_incidentes (cliente_instalacion, estado, user_email, creado_en, firma_responsable) "
        "VALUES ('Instalación B1', 'Abierto', %s, NOW(), NULL)", [ADMIN])
    # La planilla vehicular ya no expone su firma (columna obsoleta en el
    # detalle): entra al backup sin imagen. La supervisión mapea dos firmas.
    sql("INSERT INTO planilla_vehicular (cliente_instalacion, submitted_by_email, creado_en, placa_vehiculo, firma_responsable) "
        "VALUES ('Instalación B1', %s, NOW(), 'BK-001', %s)", [ADMIN, FIRMA])
    sql("INSERT INTO supervision_puesto (cliente_instalacion, supervisor, submitted_by_email, creado_en, "
        "firma_supervisor, firma_guardia) VALUES ('Instalación B1', 'Sup Prueba', %s, NOW(), %s, %s)",
        [ADMIN, FIRMA, FIRMA])


class BackupTests(unittest.TestCase):
    admin = None

    @classmethod
    def setUpClass(cls):
        cls.admin = A.app.test_client()
        assert login(cls.admin, ADMIN).status_code == 302

    def _backup(self, body=PERIODO):
        return self.admin.post('/viewer/api/backup', json=body)

    def test_01_el_zip_llega_por_partes_y_abre(self):
        r = self._backup()
        self.assertEqual(r.status_code, 200, r.data[:300])
        self.assertEqual(r.mimetype, 'application/zip')
        self.assertNotIn('Content-Length', r.headers, 'sin tamaño declarado: respuesta por partes')
        self.assertIn('attachment; filename="backup_Todos_los_clientes_2000-01-01_2099-12-31_',
                      r.headers.get('Content-Disposition', ''))
        z = zipfile.ZipFile(io.BytesIO(r.data))
        self.assertIsNone(z.testzip())
        nombres = z.namelist()
        self.assertIn('manifiesto.json', nombres)
        self.assertIn('datos/reporte_incidente.json', nombres)
        self.assertIn('datos/planilla_vehicular.json', nombres)
        self.assertIn('datos/supervision_puesto.json', nombres)
        self.assertEqual(len([n for n in nombres if n.endswith('.xlsx')]), 1)
        m = json.loads(z.read('manifiesto.json'))
        self.assertEqual(m['total_registros'], 4)
        self.assertEqual(m['registros_por_formulario'],
                         {'planilla_vehicular': 1, 'reporte_incidente': 2, 'supervision_puesto': 1})
        self.assertEqual(m['firmas_extraidas'], 3, 'una del incidente y dos de la supervisión')
        self.assertEqual(len([n for n in nombres if n.startswith('firmas/supervision_puesto/')]), 2)

    def test_02_las_firmas_van_como_archivo_y_el_registro_guarda_la_ruta(self):
        z = zipfile.ZipFile(io.BytesIO(self._backup().data))
        filas = json.loads(z.read('datos/reporte_incidente.json'))
        con_firma = [f for f in filas if any(str(v).startswith('firmas/') for v in f['data'].values())]
        self.assertEqual(len(con_firma), 1, 'sólo un incidente tiene firma')
        ruta = next(v for v in con_firma[0]['data'].values() if str(v).startswith('firmas/'))
        self.assertTrue(ruta.startswith('firmas/reporte_incidente/') and ruta.endswith('.png'), ruta)
        self.assertEqual(z.read(ruta), PNG, 'la imagen del ZIP es la firma original, byte a byte')
        for f in ('reporte_incidente', 'planilla_vehicular', 'supervision_puesto'):
            self.assertNotIn('data:image', z.read(f'datos/{f}.json').decode('utf-8'), f)
        sup = json.loads(z.read('datos/supervision_puesto.json'))[0]['data']
        self.assertTrue(sup['Firma Supervisor'].startswith('firmas/supervision_puesto/'))
        self.assertTrue(sup['Firma del Guardia'].startswith('firmas/supervision_puesto/'))
        self.assertNotEqual(sup['Firma Supervisor'], sup['Firma del Guardia'])

    def test_03_el_excel_lleva_la_ruta_y_no_el_base64(self):
        from openpyxl import load_workbook
        z = zipfile.ZipFile(io.BytesIO(self._backup().data))
        xlsx = next(n for n in z.namelist() if n.endswith('.xlsx'))
        wb = load_workbook(io.BytesIO(z.read(xlsx)), read_only=True)
        celdas = [str(c) for ws in wb.worksheets for fila in ws.iter_rows(values_only=True) for c in fila if c]
        self.assertTrue(any(c.startswith('firmas/supervision_puesto/') for c in celdas))
        self.assertFalse(any('data:image' in c for c in celdas))
        self.assertEqual(sorted(ws.title for ws in wb.worksheets),
                         ['Control de Supervisión', 'Pre-Operacional Vehicular', 'Reporte de Incidente'])

    def test_04_las_firmas_anidadas_tambien_salen(self):
        import viewer_bp as V
        buf = io.BytesIO()
        contador = [0]
        with zipfile.ZipFile(buf, 'w') as z:
            datos = V._backup_extraer_imagenes(
                {'Lista de Asistencia': [{'nombre': 'Ana', 'firma': FIRMA}, {'nombre': 'Luis', 'firma': ''}],
                 'Firma': FIRMA, 'Nota': 'data:image/png;base64,%%no-es-base64%%', 'Total': 3},
                'capacitacion', [7], z, contador)
        self.assertEqual(contador[0], 2)
        self.assertEqual(datos['Lista de Asistencia'][0]['firma'], 'firmas/capacitacion/7_Lista_de_Asistencia_0_firma_1.png')
        self.assertEqual(datos['Lista de Asistencia'][1]['firma'], '')
        self.assertEqual(datos['Firma'], 'firmas/capacitacion/7_Firma_2.png')
        self.assertEqual(datos['Total'], 3)
        self.assertTrue(datos['Nota'].startswith('data:image'), 'lo que no decodifica se deja intacto')
        self.assertEqual(zipfile.ZipFile(buf).read('firmas/capacitacion/7_Firma_2.png'), PNG)

    def test_05_se_registra_solo_si_el_zip_se_transmitio_completo(self):
        sql("DELETE FROM backups_realizados")
        self.assertEqual(self._backup({'start_date': '2099-01-01', 'end_date': '2000-01-01'}).status_code, 400)
        self.assertEqual(sql("SELECT COUNT(*) AS n FROM backups_realizados", uno=True)['n'], 0)

        import viewer_bp as V
        original = V.fetch_reports

        def falla(*a, **k):
            raise RuntimeError('base caída a mitad del backup')
        V.fetch_reports = falla
        try:
            r = self._backup()
            self.assertEqual(r.status_code, 200, 'ya se estaba transmitiendo: el código no puede cambiar')
            r.data
        finally:
            V.fetch_reports = original
        self.assertEqual(sql("SELECT COUNT(*) AS n FROM backups_realizados", uno=True)['n'], 0,
                         'un ZIP truncado no cuenta como backup')

        r = self._backup()
        self.assertIsNone(zipfile.ZipFile(io.BytesIO(r.data)).testzip())
        fila = sql("SELECT generado_por, total_registros, formato, cliente_nombre FROM backups_realizados", uno=True)
        self.assertEqual((fila['generado_por'], fila['total_registros'], fila['formato'], fila['cliente_nombre']),
                         (ADMIN, 4, 'zip', None))
        e = self.admin.get('/viewer/api/backup/estado').get_json()
        self.assertEqual(e['ultimo']['total_registros'], 4)


if __name__ == '__main__':
    unittest.main(verbosity=2)
