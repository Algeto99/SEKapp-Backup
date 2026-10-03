"""Perfil Coordinador (coordinador.py): el Coordinador ve y gestiona sólo su ámbito.

Usa el arranque común de tests/sekapp_testing.py (Postgres desechable, esquema
recreado desde sql/schema.sql). Correr con:
    monolith/venv/bin/python tests/test_coordinador.py
"""
import unittest
from datetime import date, timedelta
from io import BytesIO

from sekapp_testing import (A, sql, eventos, configurar_app, recrear_base, crear_empresa,
                            crear_usuario, login)

ADMIN = 'admin@pruebas.sekapp'
COORD = 'coordinador@pruebas.sekapp'
COORD2 = 'coordinador2@pruebas.sekapp'      # sin ámbito
SUP = 'supervisor@pruebas.sekapp'
D = {}   # ids sembrados


def setUpModule():
    configurar_app()
    recrear_base()
    cid = crear_empresa()
    D['admin'] = crear_usuario(ADMIN, 'Admin Pruebas', cid, is_admin=True, is_super_admin=True)
    D['coord'] = crear_usuario(COORD, 'Coordinador Pruebas', cid, is_coordinador=True)
    D['coord2'] = crear_usuario(COORD2, 'Coordinador Sin Ámbito', cid, is_coordinador=True)
    D['sup'] = crear_usuario(SUP, 'Supervisor Pruebas', cid)

    # Dos clientes con una instalación cada uno.
    for letra in ('A', 'B'):
        c = sql("INSERT INTO customer_companies (company_id, name, code, is_active) "
                "VALUES (%s, %s, %s, TRUE) RETURNING id", [cid, f'Cliente {letra}', letra], uno=True)['id']
        p = sql("INSERT INTO propiedades (nombre, activa, customer_company_id) "
                "VALUES (%s, TRUE, %s) RETURNING id_propiedad", [f'Instalación {letra}1', c], uno=True)['id_propiedad']
        D[f'cli{letra}'], D[f'prop{letra}'] = c, p
        # Incidente abierto hace 3 días (regla 2) y puesto sin supervisión hace 5 días (regla 1).
        D[f'inc{letra}'] = sql(
            "INSERT INTO reportes_incidentes (cliente_instalacion, id_propiedad, customer_company_id, company_id, "
            "estado, tipo_incidente, user_email, fecha_hora, creado_en) "
            "VALUES (%s, %s, %s, %s, 'Reportado', 'Prueba', %s, NOW() - INTERVAL '3 days', NOW() - INTERVAL '3 days') "
            "RETURNING id_reporte_incidente",
            [f'Instalación {letra}1', p, c, cid, SUP], uno=True)['id_reporte_incidente']
        D[f'sup{letra}'] = sql(
            "INSERT INTO supervision_puesto (cliente_instalacion, id_propiedad, customer_company_id, company_id, "
            "supervisor, submitted_by_email, fecha_hora, creado_en) "
            "VALUES (%s, %s, %s, %s, 'Sup Prueba', %s, NOW() - INTERVAL '5 days', NOW() - INTERVAL '5 days') "
            "RETURNING id_supervision",
            [f'Instalación {letra}1', p, c, cid, SUP], uno=True)['id_supervision']

    # Ámbito del Coordinador: el Cliente A completo.
    sql("INSERT INTO coordinador_ambito (user_id, customer_company_id, creado_por) VALUES (%s, %s, %s)",
        [D['coord'], D['cliA'], ADMIN])


class CoordinadorTests(unittest.TestCase):
    admin = coord = coord2 = sup = None

    @classmethod
    def setUpClass(cls):
        cls.admin, cls.coord, cls.coord2, cls.sup = (A.app.test_client() for _ in range(4))
        for cliente, email in ((cls.admin, ADMIN), (cls.coord, COORD), (cls.coord2, COORD2), (cls.sup, SUP)):
            r = login(cliente, email)
            assert r.status_code == 302, (email, r.status_code)
        # Asignaciones de hallazgo sobre los dos incidentes, al Supervisor (crea las columnas de la tabla).
        for letra in ('A', 'B'):
            r = cls.admin.post('/cgeo/api/asignar-hallazgo',
                               json={'form_type': 'reporte_incidente', 'record_id': D[f'inc{letra}'],
                                     'asignado_a': D['sup'], 'nota': f'Hallazgo {letra}'})
            assert r.status_code == 200, r.data[:200]
            D[f'asig{letra}'] = r.get_json().get('id') or sql(
                "SELECT id FROM asignaciones_hallazgo WHERE record_id = %s", [D[f'inc{letra}']], uno=True)['id']

    # ------------------------------------------------------------------
    def alertas(self, cliente, query=''):
        r = cliente.get('/cgeo/api/alertas' + query)
        self.assertEqual(r.status_code, 200, r.data[:200])
        return r.get_json()['alertas']

    @staticmethod
    def registros(alertas, form_type):
        return {a['record_id'] for a in alertas if a.get('form_type') == form_type and a.get('record_id')}

    def test_01_entrada_y_rol(self):
        c = A.app.test_client()          # otro dispositivo: el Coordinador ya tiene sesión en self.coord
        r = login(c, COORD)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'USUARIO YA ACTIVO', r.data.upper())
        r = c.post('/login', data={'confirmar_sesion': '1'})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r.headers['Location'].endswith('/matrices/alertas'), r.headers['Location'])
        self.assertEqual(c.get('/landing/user_info').get_json()['roles'], ['coordinador'])
        self.assertIn('Coordinador', c.get('/landing/').get_data(as_text=True))
        # con sesión vigente, GET /login también lo manda a Matrices
        self.assertTrue(c.get('/login').headers['Location'].endswith('/matrices/alertas'))
        e = eventos(accion='Inicio de sesión', usuario_email=COORD)[-1]
        self.assertTrue(e['detalle']['es_coordinador'])
        c.get('/logout')

    def test_02_restricciones_de_navegacion(self):
        for ruta in ('/cgeo/morning-briefing/', '/cgeo/recursos/', '/cgeo/operacion/', '/dashboard/',
                     '/dashboard/supervision/', '/expediente/', '/admin/', '/admin/auditoria', '/admin/thresholds'):
            r = self.coord.get(ruta)
            self.assertIn(r.status_code, (302, 403), f'{ruta} respondió {r.status_code}')
        self.assertEqual(self.coord.get('/cgeo/api/morning-briefing-data').status_code, 403)
        self.assertEqual(self.coord.get('/admin/api/auditoria').status_code, 403)
        # lo que sí puede abrir
        for ruta in ('/landing/', '/forms/select', '/matrices/', '/matrices/alertas', '/cgeo/hallazgos',
                     '/dashboard/incidentes/', '/dashboard/visitas/'):
            self.assertEqual(self.coord.get(ruta).status_code, 200, ruta)

    def test_03_pantalla_matrices(self):
        html = self.coord.get('/matrices/alertas').get_data(as_text=True)
        self.assertIn('Cliente A', html)
        self.assertIn('su ámbito', html)
        self.assertEqual(self.sup.get('/matrices/alertas').status_code, 302, 'el Supervisor no entra')
        self.assertEqual(self.admin.get('/matrices/alertas').status_code, 200)
        self.assertIn('id="alertasCard"', self.coord.get('/matrices/').get_data(as_text=True))
        self.assertNotIn('id="alertasCard"', self.sup.get('/matrices/').get_data(as_text=True))
        self.assertTrue(eventos(accion='Consulta de Alertas / Novedades', usuario_email=COORD))

    def test_04_alertas_acotadas(self):
        a_coord = self.alertas(self.coord)
        a_admin = self.alertas(self.admin)
        self.assertEqual(self.registros(a_coord, 'reporte_incidente'), {D['incA']})
        self.assertEqual(self.registros(a_admin, 'reporte_incidente'), {D['incA'], D['incB']})
        self.assertEqual(self.registros(a_coord, 'supervision_puesto'), {D['supA']})
        textos = ' '.join(a['texto'] for a in a_coord)
        self.assertNotIn('B1', textos)
        self.assertIn('Instalación A1', textos)
        # asignaciones pendientes: sólo la del incidente A
        asig = {a['asignacion_id'] for a in a_coord if a.get('asignacion_id')}
        self.assertEqual(asig, {D['asigA']})
        # reglas excluidas para el Coordinador
        self.assertFalse([a for a in a_coord if a.get('regla') in (13, 15)])
        # el Supervisor sigue sin acceso al endpoint de alertas
        self.assertEqual(self.sup.get('/cgeo/api/alertas').status_code, 403)

    def test_05_filtros_fuera_del_ambito_se_ignoran(self):
        con_b = self.alertas(self.coord, f"?cliente={D['cliB']}")
        self.assertEqual(self.registros(con_b, 'reporte_incidente'), {D['incA']})
        con_b = self.alertas(self.coord, f"?propiedad={D['propB']}")
        self.assertEqual(self.registros(con_b, 'reporte_incidente'), {D['incA']})
        con_a = self.alertas(self.coord, f"?cliente={D['cliA']}")
        self.assertEqual(self.registros(con_a, 'reporte_incidente'), {D['incA']})
        f = self.coord.get('/cgeo/api/filtros').get_json()
        self.assertEqual([c['id'] for c in f['clientes']], [D['cliA']])
        self.assertEqual([p['id'] for p in f['propiedades']], [D['propA']])
        self.assertGreaterEqual(len(self.admin.get('/cgeo/api/filtros').get_json()['clientes']), 2)

    def test_06_detalle_y_exportaciones(self):
        self.assertEqual(self.coord.get(f"/viewer/api/report/{D['incA']}?form_type=reporte_incidente").status_code, 200)
        r = self.coord.get(f"/viewer/api/report/{D['incB']}?form_type=reporte_incidente")
        self.assertEqual(r.status_code, 403)
        e = eventos(accion='Visualización de detalle', usuario_email=COORD)[-1]
        self.assertEqual((e['estado'], e['registro_id']), ('Rechazado', D['incB']))
        self.assertEqual(self.coord.post('/viewer/api/export-excel',
                         json={'reports': [{'id': D['incB'], 'formType': 'reporte_incidente'}]}).status_code, 403)
        r = self.coord.post('/viewer/api/export-excel',
                            json={'reports': [{'id': D['incA'], 'formType': 'reporte_incidente'},
                                              {'id': D['incB'], 'formType': 'reporte_incidente'}]})
        self.assertEqual(r.status_code, 200, r.data[:200])
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(r.data)).active
        self.assertEqual(ws.max_row - 1, 1, 'sólo el incidente del ámbito va al Excel')
        # el Administrador exporta ambos
        r = self.admin.post('/viewer/api/export-excel',
                            json={'reports': [{'id': D['incA'], 'formType': 'reporte_incidente'},
                                              {'id': D['incB'], 'formType': 'reporte_incidente'}]})
        self.assertEqual(load_workbook(BytesIO(r.data)).active.max_row - 1, 2)

    def test_07_estado_e_historial(self):
        r = self.coord.put(f"/dashboard/api/incidentes/{D['incA']}/estado", json={'estado': 'En proceso', 'accion_tomada': 'Revisado'})
        self.assertEqual(r.status_code, 200, r.data[:200])
        self.assertEqual(sql("SELECT estado FROM reportes_incidentes WHERE id_reporte_incidente = %s", [D['incA']], uno=True)['estado'], 'En proceso')
        r = self.coord.put(f"/dashboard/api/incidentes/{D['incB']}/estado", json={'estado': 'En proceso', 'accion_tomada': 'x'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(sql("SELECT estado FROM reportes_incidentes WHERE id_reporte_incidente = %s", [D['incB']], uno=True)['estado'],
                         'Asignado', 'el incidente B no cambió (quedó Asignado al crear el hallazgo)')
        self.assertEqual(eventos(accion='Cambio de estado', usuario_email=COORD)[-1]['estado'], 'Rechazado')
        self.assertEqual(self.coord.get(f"/dashboard/api/incidentes/{D['incA']}/historial").status_code, 200)
        self.assertEqual(self.coord.get(f"/dashboard/api/incidentes/{D['incB']}/historial").status_code, 403)

    def test_08_hallazgos(self):
        r = self.coord.post('/cgeo/api/asignar-hallazgo',
                            json={'form_type': 'reporte_incidente', 'record_id': D['incB'], 'asignado_a': D['sup']})
        self.assertEqual(r.status_code, 403)
        r = self.coord.post('/cgeo/api/asignar-hallazgo',
                            json={'form_type': 'reporte_incidente', 'record_id': D['incA'], 'asignado_a': D['sup'],
                                  'nota': 'Reasignado por el Coordinador'})
        self.assertEqual(r.status_code, 200, r.data[:200])
        self.assertEqual(self.coord.get(f"/cgeo/hallazgo/{D['asigB']}").status_code, 403)
        self.assertEqual(self.coord.get(f"/cgeo/hallazgo/{D['asigA']}").status_code, 200)
        # El enlace de retorno depende del perfil y de la procedencia: el Coordinador
        # y el Supervisor vuelven siempre a Hallazgos Asignados (no tienen briefing),
        # aunque el navegador diga que vinieron de él; el Administrador vuelve al
        # briefing sólo cuando viene de ahí, conservando su query.
        url = f"/cgeo/hallazgo/{D['asigA']}"
        briefing = {'Referer': 'http://localhost/cgeo/morning-briefing/?cliente=1'}
        for quien in (self.coord, self.sup):
            html = quien.get(url, headers=briefing).get_data(as_text=True)
            self.assertIn('Volver a Hallazgos Asignados', html)
            self.assertNotIn('Morning Briefing', html)
        self.assertIn('Volver a Hallazgos Asignados', self.admin.get(url).get_data(as_text=True))
        html = self.admin.get(url, headers=briefing).get_data(as_text=True)
        self.assertIn('href="/cgeo/morning-briefing/?cliente=1"', html)
        self.assertIn('Volver al Morning Briefing', html)
        html = self.admin.get(url, headers={'Referer': 'http://localhost/cgeo/operacion/'}).get_data(as_text=True)
        self.assertIn('Volver a Operación e Incidentes', html)
        html = self.admin.get(url, headers={'Referer': 'http://localhost/cgeo/hallazgos'}).get_data(as_text=True)
        self.assertIn('Volver a Hallazgos Asignados', html)
        self.assertEqual(self.coord.post(f"/cgeo/api/asignaciones/{D['asigB']}/gestionar", json={'nota': 'x'}).status_code, 403)
        r = self.coord.post(f"/cgeo/api/asignaciones/{D['asigA']}/gestionar", json={'nota': 'Atendido'})
        self.assertEqual(r.status_code, 200, r.data[:200])
        fila = sql("SELECT estado, cerrado_por FROM asignaciones_hallazgo WHERE id = %s", [D['asigA']], uno=True)
        self.assertEqual((fila['estado'], fila['cerrado_por']), ('Gestionado', COORD))
        self.assertTrue(eventos(accion='Gestión de hallazgo', usuario_email=COORD))

    def test_09_matrices_acotadas(self):
        ids = lambda cliente: {r.get('id') or r.get('id_reporte_incidente') for r in
                               cliente.get('/dashboard/api/incidentes/detalles').get_json()['detalles']}
        self.assertEqual(ids(self.coord), {D['incA']})
        self.assertEqual(ids(self.admin), {D['incA'], D['incB']})
        self.assertEqual(ids(self.sup), {D['incA'], D['incB']}, 'el Supervisor no cambia')
        self.assertEqual(self.coord.get('/dashboard/api/incidentes/clientes').get_json()['clientes'], ['Instalación A1'])
        # Sin rango, el hub usa el mes en curso y el incidente (hace 3 días) queda
        # fuera los primeros días de cada mes: se pide un rango explícito.
        hoy = date.today()
        rango = f"?date_from={hoy - timedelta(days=30):%Y-%m-%d}&date_to={hoy:%Y-%m-%d}"
        self.assertEqual(self.coord.get('/matrices/api/stats' + rango).get_json()['incidentes']['total'], 1)
        self.assertEqual(self.admin.get('/matrices/api/stats' + rango).get_json()['incidentes']['total'], 2)

    def test_09b_supervisiones_por_periodo(self):
        # Alertas / Novedades → Supervisiones + Período: la lista de supervisiones
        # registradas sale del endpoint de detalles, acotado al ámbito, entre
        # `desde` y `hasta` (inclusive por fecha).
        def ids(cliente, query):
            r = cliente.get('/dashboard/api/supervision/detalles' + query)
            self.assertEqual(r.status_code, 200, r.data[:200])
            return {d['id'] for d in r.get_json()['detalles']}
        hoy = date.today()
        d30 = f"{hoy - timedelta(days=30):%Y-%m-%d}"
        self.assertEqual(ids(self.coord, f"?desde={d30}&hasta={hoy:%Y-%m-%d}"), {D['supA']})
        self.assertEqual(ids(self.admin, f"?desde={d30}&hasta={hoy:%Y-%m-%d}"), {D['supA'], D['supB']})
        self.assertEqual(ids(self.coord2, f"?desde={d30}&hasta={hoy:%Y-%m-%d}"), set(), 'sin ámbito no ve nada')
        # `hasta` recorta: la supervisión sembrada es de hace 5 días.
        self.assertEqual(ids(self.admin, f"?desde={d30}&hasta={hoy - timedelta(days=6):%Y-%m-%d}"), set())
        # Un `hasta` mal formado se ignora en vez de romper la consulta.
        self.assertEqual(ids(self.coord, '?hasta=ayer'), {D['supA']})

    def test_10_coordinador_sin_ambito(self):
        self.assertEqual(self.alertas(self.coord2), [])
        self.assertEqual(self.coord2.get(f"/viewer/api/report/{D['incA']}?form_type=reporte_incidente").status_code, 403)
        self.assertEqual(self.coord2.get('/dashboard/api/incidentes/detalles').get_json()['detalles'], [])
        self.assertIn('No tiene un ámbito', self.coord2.get('/matrices/alertas').get_data(as_text=True))

    def test_11_editor_de_ambito(self):
        html = self.admin.get(f"/admin/users/{D['coord']}/ambito").get_data(as_text=True)
        self.assertIn('Cliente A', html)
        self.assertIn('Instalación B1', html)
        self.assertEqual(self.coord.get(f"/admin/users/{D['coord']}/ambito").status_code, 302)
        # sólo la instalación B1: el Coordinador pasa a ver B y deja de ver A
        r = self.admin.post(f"/admin/users/{D['coord']}/ambito", data={'propiedades': [str(D['propB'])]})
        self.assertEqual(r.status_code, 302)
        filas = sql("SELECT customer_company_id, id_propiedad FROM coordinador_ambito WHERE user_id = %s", [D['coord']])
        self.assertEqual([(f['customer_company_id'], f['id_propiedad']) for f in filas], [(None, D['propB'])])
        a = self.alertas(self.coord)
        self.assertEqual(self.registros(a, 'reporte_incidente'), {D['incB']})
        self.assertEqual(self.coord.get(f"/viewer/api/report/{D['incA']}?form_type=reporte_incidente").status_code, 403)
        f = self.coord.get('/cgeo/api/filtros').get_json()
        self.assertEqual([c['id'] for c in f['clientes']], [D['cliB']])
        e = eventos(accion='Modificación de ámbito de coordinador', usuario_email=ADMIN)[-1]
        self.assertEqual(e['registro_id'], D['coord'])
        self.assertEqual(e['detalle']['propiedades'], str(D['propB']))
        # panel: distintivo y conteo
        panel = self.admin.get('/admin/').get_data(as_text=True)
        self.assertIn('Coordinador', panel)
        self.assertIn('Ámbito: 1', panel)

    def test_12_quitar_rol(self):
        r = self.admin.post(f"/admin/users/{D['coord']}/toggle-coordinador")
        self.assertEqual(r.status_code, 302)
        self.assertFalse(sql("SELECT is_coordinador FROM users WHERE id = %s", [D['coord']], uno=True)['is_coordinador'])
        # el JWT vigente aún dice Coordinador, pero la verificación en base ya lo rechaza
        self.assertEqual(self.coord.get('/matrices/alertas').status_code, 302)
        self.assertEqual(self.coord.get('/cgeo/api/alertas').status_code, 403)
        self.assertTrue(eventos(accion='Cambio de rol de coordinador', usuario_email=ADMIN))
        # volver a darle el rol manda al editor de ámbito
        r = self.admin.post(f"/admin/users/{D['coord']}/toggle-coordinador")
        self.assertTrue(r.headers['Location'].endswith(f"/admin/users/{D['coord']}/ambito"))


if __name__ == '__main__':
    unittest.main(verbosity=2)
