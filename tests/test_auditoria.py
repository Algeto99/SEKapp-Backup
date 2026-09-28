"""Log de eventos (auditoria.py): prueba de extremo a extremo con el cliente de Flask.

Usa el arranque común de tests/sekapp_testing.py (Postgres desechable, esquema
recreado desde sql/schema.sql). Correr con:
    monolith/venv/bin/python tests/test_auditoria.py
"""
import json
import unittest
import zoneinfo
from datetime import datetime, timedelta
from io import BytesIO

import psycopg2

from sekapp_testing import (A, RAIZ, CLAVE, ZONA, auditoria, conectar, sql, eventos,
                            configurar_app, recrear_base, crear_empresa, crear_usuario, login)

ADMIN = 'admin@pruebas.sekapp'
SUP = 'supervisor@pruebas.sekapp'


def setUpModule():
    configurar_app()
    recrear_base()
    cid = crear_empresa()
    crear_usuario(ADMIN, 'Admin Pruebas', cid, is_admin=True, is_super_admin=True)
    crear_usuario(SUP, 'Supervisor Pruebas', cid)


class AuditoriaTests(unittest.TestCase):
    admin = None
    sup = None
    estado = {}

    @classmethod
    def setUpClass(cls):
        cls.admin = A.app.test_client()
        cls.sup = A.app.test_client()

    # ------------------------------------------------------------------
    def test_01_ruido_no_genera_eventos(self):
        c = A.app.test_client()
        for ruta in ('/health', '/favicon.ico', '/noexiste', '/static/css/forms.css', '/forms/api/csrf_token', '/login'):
            c.get(ruta)
        self.assertEqual(eventos(), [], 'ninguna de esas rutas debe registrarse')

    def test_02_login_fallido_y_exitoso(self):
        r = login(self.sup, SUP, 'incorrecta')
        self.assertEqual(r.status_code, 200)
        ev = eventos(accion='Inicio de sesión fallido')
        self.assertEqual(len(ev), 1)
        self.assertEqual((ev[0]['usuario_email'], ev[0]['estado'], ev[0]['tipo_evento']), (SUP, 'Rechazado', 'sesion'))
        self.assertEqual(ev[0]['detalle']['motivo'], 'contraseña incorrecta')

        login(self.sup, 'nadie@pruebas.sekapp', 'x')
        self.assertEqual(eventos(accion='Inicio de sesión fallido')[-1]['detalle']['motivo'], 'usuario no existe')

        r = login(self.sup, SUP, CLAVE)
        self.assertEqual(r.status_code, 302)
        ev = eventos(accion='Inicio de sesión', usuario_email=SUP)
        self.assertEqual(len(ev), 1)
        e = ev[0]
        self.assertEqual(e['estado'], 'Exitoso')
        self.assertEqual(e['usuario_nombre'], 'Supervisor Pruebas')
        self.assertEqual(e['licencia'], 'Kanan Sentinel Pruebas')
        self.assertTrue(e['sesion_jti'])
        self.assertEqual(e['dispositivo'], 'Computador')
        self.assertEqual(e['modulo'], 'Acceso')
        self.assertEqual(e['metodo_ruta'], 'POST /login')
        self.assertTrue(e['ip'])
        # el jti coincide con la sesión registrada
        ses = sql("SELECT jti FROM sesiones_usuario WHERE usuario_email = %s", [SUP])
        self.assertEqual(ses[0]['jti'], e['sesion_jti'])

        r = login(self.admin, ADMIN, CLAVE)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(len(eventos(accion='Inicio de sesión', usuario_email=ADMIN)), 1)

    def test_03_aviso_de_sesion_activa(self):
        otro = A.app.test_client()          # otro dispositivo, mismas credenciales
        r = login(otro, ADMIN, CLAVE)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'USUARIO YA ACTIVO', r.data.upper())
        ev = eventos(accion='Aviso de sesión activa')
        self.assertEqual(len(ev), 1)
        self.assertEqual((ev[0]['usuario_email'], ev[0]['estado']), (ADMIN, 'Pendiente'))
        r = otro.post('/login', data={'confirmar_sesion': '1'})
        self.assertEqual(r.status_code, 302)
        e = eventos(accion='Inicio de sesión', usuario_email=ADMIN)[-1]
        self.assertTrue(e['detalle']['confirmo_sesion_activa'])
        otro.get('/logout')

    def test_04_consultas_de_modulo(self):
        self.admin.get('/landing/')
        self.assertEqual(eventos(accion='Acceso a la plataforma', usuario_email=ADMIN)[0]['modulo'], 'Inicio')
        self.admin.get('/dashboard/incidentes/')
        self.assertEqual(eventos(modulo='Dashboards: Incidentes', accion='Consulta de módulo')[0]['usuario_email'], ADMIN)
        antes = len(eventos())
        self.admin.get('/dashboard/api/incidentes/data')
        self.admin.get('/dashboard/api/incidentes/filtros')
        self.assertEqual(len(eventos()), antes, 'las llamadas de datos y filtros no se registran')
        self.admin.get('/dashboard/api/incidentes/detalles?cliente=Cliente%20X&year=2026')
        e = eventos(accion='Visualización de detalle', modulo='Dashboards: Incidentes')[0]
        self.assertEqual(e['detalle']['filtros'], {'cliente': 'Cliente X', 'year': '2026'})
        self.assertEqual(e['formulario'], 'reporte_incidente')
        self.admin.get('/viewer/')
        self.assertTrue(eventos(modulo='Reportes', accion='Consulta del módulo Reportes'))
        self.sup.get('/forms/select')
        self.assertTrue(eventos(modulo='Formularios', accion='Consulta del módulo Formularios', usuario_email=SUP))

    def test_05_envio_y_edicion_de_formulario(self):
        datos = {
            'cliente_instalacion': 'Cliente Prueba', 'fecha_hora': '2026-09-27T10:00',
            'rol_aplicador': 'Supervisor de Seguridad', 'nombre_responsable': 'Sup Prueba',
            'atencion_cliente': '5', 'comunicacion': '4', 'confiabilidad': '5', 'capacidad_reaccion': '4',
            'cumplimiento': '5', 'competencia_personal': '5', 'actitud_servicio': '4', 'atencion_quejas': '5',
            'calificacion_global_nps': '9', 'recomendaria_servicio': 'Sí', 'encuestado': 'Cliente',
            'submitter_timezone': ZONA,
        }
        self.sup.get('/forms/medicion_experiencia_cliente')
        self.assertTrue(eventos(accion='Apertura de formulario', formulario='medicion_experiencia_cliente'))
        r = self.sup.post('/forms/submit_medicion_experiencia_cliente', data=datos,
                          headers={'X-Requested-With': 'XMLHttpRequest'})
        self.assertEqual(r.status_code, 200, r.data[:300])
        fila = sql("SELECT id_encuesta, submitted_by_email FROM medicion_experiencia_cliente ORDER BY id_encuesta DESC LIMIT 1", uno=True)
        self.assertIsNotNone(fila)
        ev = eventos(accion='Envío de formulario', formulario='medicion_experiencia_cliente')
        self.assertEqual(len(ev), 1)
        self.assertEqual(ev[0]['registro_id'], fila['id_encuesta'])
        self.assertEqual((ev[0]['usuario_email'], ev[0]['tipo_evento'], ev[0]['estado']), (SUP, 'envio', 'Exitoso'))
        AuditoriaTests.estado['encuesta'] = fila['id_encuesta']

        # edición (ruta de administrador) con motivo
        eid = fila['id_encuesta']
        self.admin.get(f'/forms/medicion_experiencia_cliente/{eid}/editar')
        self.assertEqual(eventos(accion='Apertura de registro para edición')[0]['registro_id'], eid)
        datos_ed = dict(datos, comunicacion='2', motivo='Corrección de información')
        r = self.admin.post(f'/forms/submit_medicion_experiencia_cliente/{eid}/editar', data=datos_ed,
                            headers={'X-Requested-With': 'XMLHttpRequest'})
        self.assertEqual(r.status_code, 200, r.data[:300])
        ev = eventos(accion='Edición de registro', formulario='medicion_experiencia_cliente')
        self.assertEqual(len(ev), 1)
        self.assertEqual((ev[0]['registro_id'], ev[0]['usuario_email'], ev[0]['tipo_evento']), (eid, ADMIN, 'edicion'))
        self.assertEqual(ev[0]['detalle']['motivo'], 'Corrección de información')
        self.assertTrue(sql("SELECT 1 AS x FROM formulario_edicion_historial WHERE registro_id = %s", [eid]))

        # envío rechazado por validación (fecha futura): queda con estado Rechazado
        r = self.sup.post('/forms/submit_medicion_experiencia_cliente', data=dict(datos, fecha_hora='2030-01-01T10:00'),
                          headers={'X-Requested-With': 'XMLHttpRequest'})
        self.assertEqual(r.status_code, 400)
        e = eventos(accion='Envío de formulario', formulario='medicion_experiencia_cliente')[-1]
        self.assertEqual((e['estado'], e['registro_id']), ('Rechazado', None))

    def test_06_detalle_y_exportacion(self):
        eid = self.estado['encuesta']
        r = self.admin.get(f'/viewer/api/report/{eid}?form_type=medicion_experiencia_cliente')
        self.assertEqual(r.status_code, 200)
        e = eventos(modulo='Reportes', accion='Visualización de detalle')[0]
        self.assertEqual((e['registro_id'], e['formulario']), (eid, 'medicion_experiencia_cliente'))
        r = self.admin.post('/viewer/api/export-excel',
                            json={'reports': [{'id': eid, 'formType': 'medicion_experiencia_cliente'}]})
        self.assertEqual(r.status_code, 200, r.data[:300])
        e = eventos(accion='Exportación a Excel')[0]
        self.assertEqual(e['tipo_evento'], 'descarga')
        self.assertEqual(e['detalle']['reports'], [{'id': eid, 'formType': 'medicion_experiencia_cliente'}])
        # generación de PDF: sin WeasyPrint local responde 503 y el evento queda como Error
        r = self.admin.post('/viewer/api/generate-pdf',
                            json={'reports': [{'id': eid, 'formType': 'medicion_experiencia_cliente'}]})
        e = eventos(accion='Generación de PDF')[0]
        self.assertEqual(e['http_status'], r.status_code)
        self.assertEqual(e['estado'], 'Exitoso' if r.status_code < 400 else 'Error')

    def test_07_cambio_de_estado_de_incidente(self):
        inc = sql("INSERT INTO reportes_incidentes (cliente_instalacion, estado, user_email, creado_en) "
                  "VALUES ('Cliente Prueba', 'Abierto', %s, NOW()) RETURNING id_reporte_incidente", [SUP], uno=True)
        iid = inc['id_reporte_incidente']
        AuditoriaTests.estado['incidente'] = iid
        r = self.admin.put(f'/dashboard/api/incidentes/{iid}/estado', json={'estado': 'Cerrado', 'accion_tomada': 'Atendido'})
        self.assertEqual(r.status_code, 200, r.data[:300])
        e = eventos(accion='Cambio de estado', formulario='reporte_incidente')[0]
        self.assertEqual((e['registro_id'], e['tipo_evento'], e['usuario_email']), (iid, 'estado', ADMIN))
        self.assertEqual(e['detalle']['estado_anterior'], 'Abierto')
        self.assertEqual(e['detalle']['estado_nuevo'], 'Cerrado')
        r = self.admin.put('/dashboard/api/incidentes/999999/estado', json={'estado': 'Cerrado'})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(eventos(accion='Cambio de estado')[-1]['estado'], 'No encontrado')

    def test_08_hallazgos(self):
        iid = self.estado['incidente']
        sup_id = sql("SELECT id FROM users WHERE email = %s", [SUP], uno=True)['id']
        r = self.admin.post('/cgeo/api/asignar-hallazgo',
                            json={'form_type': 'reporte_incidente', 'record_id': iid, 'asignado_a': sup_id,
                                  'fecha_limite': '2026-10-15', 'nota': 'Revisar', 'hallazgo_titulo': 'Prueba'})
        self.assertEqual(r.status_code, 200, r.data[:300])
        e = eventos(accion='Asignación de hallazgo')[0]
        self.assertEqual((e['formulario'], e['registro_id'], e['tipo_evento']), ('reporte_incidente', iid, 'asignacion'))
        self.assertEqual(e['detalle']['asignado_nombre'], 'Supervisor Pruebas')
        self.assertEqual(e['detalle']['hallazgo_titulo'], 'Prueba')
        asig = e['detalle']['asignacion_id']

        r = self.sup.get(f'/cgeo/hallazgo/{asig}')
        e = eventos(accion='Visualización de hallazgo')[0]
        self.assertEqual((e['registro_id'], e['formulario'], e['usuario_email']), (iid, 'reporte_incidente', SUP))
        self.assertEqual(e['detalle']['asignacion_id'], asig)
        self.assertEqual(e['http_status'], r.status_code)

        r = self.sup.post(f'/cgeo/api/asignaciones/{asig}/gestionar', json={'nota': 'Listo'})
        self.assertEqual(r.status_code, 200, r.data[:300])
        e = eventos(accion='Gestión de hallazgo')[0]
        self.assertEqual((e['registro_id'], e['formulario'], e['usuario_email']), (iid, 'reporte_incidente', SUP))
        self.assertEqual(e['detalle']['nota'], 'Listo')

    def test_09_vista_publica_qr(self):
        sup_row = sql("INSERT INTO supervision_puesto (supervisor, submitted_by_email, cliente_instalacion, creado_en) "
                      "VALUES ('Sup Prueba', %s, 'Cliente Prueba', NOW()) RETURNING id_supervision", [SUP], uno=True)
        sid = sup_row['id_supervision']
        from expediente_bp import _make_qr_token
        with A.app.app_context():
            token = _make_qr_token(sid, 1)
        r = A.app.test_client().get(f'/v/{token}')
        e = eventos(accion='Consulta pública de supervisión')[0]
        self.assertIsNone(e['usuario_email'])
        self.assertEqual(e['usuario_nombre'], 'Público (QR)')
        self.assertEqual((e['registro_id'], e['formulario'], e['tipo_evento']), (sid, 'supervision_puesto', 'publico'))
        self.assertEqual(e['http_status'], r.status_code)

    def test_10_pantalla_y_api_de_auditoria(self):
        r = self.sup.get('/admin/auditoria')
        self.assertEqual(r.status_code, 302, 'el Supervisor no ve la auditoría')
        r = self.admin.get('/admin/auditoria?usuario=' + SUP)
        self.assertEqual(r.status_code, 200)
        self.assertIn('Log de eventos', r.get_data(as_text=True))
        self.assertTrue(eventos(accion='Apertura de Auditoría', usuario_email=ADMIN))

        r = self.admin.get(f'/admin/api/auditoria?usuario={SUP}&por_pagina=200')
        self.assertEqual(r.status_code, 200)
        d = r.get_json()
        self.assertGreater(d['total'], 0)
        self.assertTrue(all(e['usuario_email'] == SUP for e in d['eventos']))
        self.assertEqual(d['total'], len(eventos(usuario_email=SUP)))
        self.assertTrue(all(e['fecha'] and e['hora'] for e in d['eventos']))

        # trazabilidad de un registro: envío, apertura para edición, edición y detalle
        eid = self.estado['encuesta']
        r = self.admin.get(f'/admin/api/auditoria?formulario=medicion_experiencia_cliente&registro_id={eid}')
        acciones = {e['accion'] for e in r.get_json()['eventos']}
        self.assertTrue({'Envío de formulario', 'Edición de registro', 'Visualización de detalle',
                         'Apertura de registro para edición'} <= acciones, acciones)

        # filtros por tipo, módulo, estado y texto libre
        self.assertTrue(all(e['tipo_evento'] == 'sesion' for e in self.admin.get('/admin/api/auditoria?tipo=sesion').get_json()['eventos']))
        self.assertTrue(all(e['estado'] == 'Rechazado' for e in self.admin.get('/admin/api/auditoria?estado=Rechazado').get_json()['eventos']))
        self.assertGreater(self.admin.get('/admin/api/auditoria?q=contrase%C3%B1a%20incorrecta').get_json()['total'], 0)
        self.assertEqual(self.admin.get('/admin/api/auditoria?modulo=No%20existe').get_json()['total'], 0)
        self.assertEqual(self.admin.get('/admin/api/auditoria?desde=no-es-fecha').status_code, 400)

        # rango de fechas de hoy en la zona de la operación
        hoy = datetime.now(zoneinfo.ZoneInfo(ZONA)).date()
        manana = hoy + timedelta(days=1)
        self.assertGreater(self.admin.get(f'/admin/api/auditoria?desde={hoy}&hasta={hoy}&tz={ZONA}').get_json()['total'], 0)
        self.assertEqual(self.admin.get(f'/admin/api/auditoria?desde={manana}&tz={ZONA}').get_json()['total'], 0)
        self.assertTrue(eventos(accion='Consulta del log de eventos', usuario_email=ADMIN))

        # paginación
        d1 = self.admin.get('/admin/api/auditoria?por_pagina=10&pagina=1').get_json()
        d2 = self.admin.get('/admin/api/auditoria?por_pagina=10&pagina=2').get_json()
        self.assertEqual(len(d1['eventos']), 10)
        self.assertNotEqual(d1['eventos'][0]['id'], d2['eventos'][0]['id'])

    def test_11_exportacion_excel(self):
        total = len(eventos(usuario_email=SUP))
        r = self.admin.get(f'/admin/api/auditoria/export?usuario={SUP}')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.content_type)
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(r.data)).active
        self.assertEqual(ws.max_row - 1, total)
        self.assertEqual(ws['C2'].value, SUP)
        self.assertTrue(eventos(accion='Exportación del log de eventos a Excel', usuario_email=ADMIN))
        self.assertEqual(self.sup.get('/admin/api/auditoria/export').status_code, 403)

    def test_12_inmutabilidad(self):
        conn = conectar()
        try:
            for sentencia in ("UPDATE eventos_auditoria SET accion = 'x' WHERE id = (SELECT MIN(id) FROM eventos_auditoria)",
                              "DELETE FROM eventos_auditoria WHERE id = (SELECT MIN(id) FROM eventos_auditoria)",
                              "TRUNCATE eventos_auditoria"):
                with self.assertRaises(psycopg2.errors.RaiseException, msg=sentencia) as ctx:
                    conn.cursor().execute(sentencia)
                self.assertIn('registro de auditoría', str(ctx.exception))
                conn.rollback()
        finally:
            conn.close()
        self.assertGreater(len(eventos()), 0)

    def test_13_fronteras_de_fecha(self):
        # 04:30 UTC del 15 = 23:30 del 14 en Bogotá: pertenece al día 14 en la zona de la operación.
        sql("INSERT INTO eventos_auditoria (fecha_hora, usuario_email, tipo_evento, modulo, accion, estado) "
            "VALUES ('2026-01-15 04:30:00+00', 'frontera@pruebas.sekapp', 'consulta', 'Prueba', 'Frontera', 'Exitoso')")
        base = '/admin/api/auditoria?usuario=frontera@pruebas.sekapp&tz=' + ZONA
        self.assertEqual(self.admin.get(base + '&desde=2026-01-14&hasta=2026-01-14').get_json()['total'], 1)
        self.assertEqual(self.admin.get(base + '&desde=2026-01-15&hasta=2026-01-15').get_json()['total'], 0)
        e = self.admin.get(base).get_json()['eventos'][0]
        self.assertEqual((e['fecha'], e['hora']), ('14/01/2026', '23:30:00'))

    def test_14_backfill_retroactivo(self):
        script = (RAIZ / 'sql' / 'backfill_eventos_auditoria.sql').read_text()
        conn = conectar()
        try:
            cur = conn.cursor()
            cur.execute(script)
            conn.commit()
            avisos = [n.strip() for n in conn.notices]
        finally:
            conn.close()
        retro = eventos(origen='retroactivo')
        self.assertGreater(len(retro), 0, avisos)
        por_accion = {}
        for e in retro:
            por_accion[e['accion']] = por_accion.get(e['accion'], 0) + 1
        sesiones = sql("SELECT COUNT(*) AS n FROM sesiones_usuario", uno=True)['n']
        self.assertEqual(por_accion.get('Inicio de sesión'), sesiones)
        self.assertGreaterEqual(por_accion.get('Cierre de sesión', 0), 1)
        envios = {(e['formulario'], e['registro_id']) for e in retro if e['accion'] == 'Envío de formulario'}
        self.assertIn(('medicion_experiencia_cliente', self.estado['encuesta']), envios)
        self.assertIn(('reporte_incidente', self.estado['incidente']), envios)
        ed = [e for e in retro if e['accion'] == 'Edición de registro']
        self.assertEqual(len(ed), 1)
        self.assertEqual(ed[0]['detalle']['motivo'], 'Corrección de información')
        self.assertIn('comunicacion', ed[0]['detalle']['campos'])
        self.assertEqual(por_accion.get('Asignación de hallazgo'), 1)
        self.assertEqual(por_accion.get('Gestión de hallazgo'), 1)
        self.assertTrue(all(e['licencia'] == 'Kanan Sentinel Pruebas' for e in retro))
        # segunda corrida: no duplica
        conn = conectar()
        try:
            conn.cursor().execute(script)
            conn.commit()
            self.assertTrue(any('ya se corrió' in n for n in conn.notices), conn.notices)
        finally:
            conn.close()
        self.assertEqual(len(eventos(origen='retroactivo')), len(retro))

    def test_15_logout(self):
        self.sup.get('/logout')
        e = eventos(accion='Cierre de sesión', usuario_email=SUP)
        self.assertEqual(len(e), 1)
        self.assertEqual(e[0]['modulo'], 'Acceso')
        self.assertTrue(sql("SELECT cerrada_en FROM sesiones_usuario WHERE usuario_email = %s", [SUP])[0]['cerrada_en'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
